"""QA for the approved 50/50 course detail and secure PDF/DOCX reading.

Uses an isolated temporary DB, local DOCX file and mocked course contents.
No external AI API or network-based document conversion is used.
"""
import io
import sys
import struct
import zlib
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from pypdf import PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app as project
from chapter_material_routes import ensure_chapter_material_schema
from syllabus_learning import syllabus_hash
from syllabus_file_preview import extract_docx_text, render_docx_preview


def _sample_word():
    sample = Document()
    sample.add_heading("ĐỀ CƯƠNG HỌC PHẦN LỊCH SỬ", 0)
    sample.add_paragraph("Chương 1. Sự ra đời của Đảng")
    sample.add_paragraph("Chương 2. Lãnh đạo kháng chiến")
    t = sample.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Chương"
    t.cell(0, 1).text = "Nội dung"
    t.cell(1, 0).text = "1"
    t.cell(1, 1).text = "<script>alert('bad')</script>"
    # Real inline PNG image: preview must not silently drop common diagrams.
    def png_chunk(kind, data):
        blob = kind + data
        return struct.pack("!I", len(data)) + blob + struct.pack("!I", zlib.crc32(blob))
    ihdr = struct.pack("!2I5B", 1, 1, 8, 2, 0, 0, 0)
    picture = (b"\x89PNG\r\n\x1a\n" + png_chunk(b"IHDR", ihdr)
               + png_chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
               + png_chunk(b"IEND", b""))
    sample.add_picture(io.BytesIO(picture))
    buffer = io.BytesIO()
    sample.save(buffer)
    return buffer.getvalue()


def run():
    with TemporaryDirectory() as tmp:
        old_db, old_upload = project.DB_PATH, project.UPLOAD_DIR
        try:
            project.DB_PATH = Path(tmp) / "course.sqlite"
            project.UPLOAD_DIR = Path(tmp) / "uploads"
            project.UPLOAD_DIR.mkdir()
            project.init_db()
            src = project.UPLOAD_DIR / "docx-sample.docx"
            src.write_bytes(_sample_word())
            text = extract_docx_text(src)
            assert "Chương 1." in text and "Chương 2." in text
            assert "Nội dung" in text and "<script>" in text
            rendered = render_docx_preview(src, "de_cuong.docx")
            assert "<table>" in rendered and "<td>" in rendered
            assert "data:image/png;base64," in rendered, "Word inline image was not preserved"
            assert "&lt;script&gt;" in rendered
            assert "<script>" not in rendered

            conn = project.db()
            stamp = project.now()
            def account(name, role):
                return conn.execute("""
                    INSERT INTO users(name,email,password_hash,role,created_at)
                    VALUES(?,?,?,?,?)
                """, (name,name+"@test.local","hash",role,stamp)).lastrowid
            admin_id = account("adminpreview","admin")
            student_id = account("studentpreview","student")
            stranger_id = account("strangerpreview","student")
            sid = conn.execute("""
                INSERT INTO subjects(user_id,name,code,credits,target_score,created_at)
                VALUES(?,'Lịch sử Đảng','HIS102',3,8,?)
            """, (student_id,stamp)).lastrowid
            unrelated=conn.execute("""
                INSERT INTO subjects(user_id,name,code,credits,created_at)
                VALUES(?,'Môn khác','DIFFERENT123',3,?)
            """, (stranger_id,stamp)).lastrowid
            conn.commit()
            admin=project.app.test_client()
            student=project.app.test_client()
            outsider=project.app.test_client()
            for client,uid in ((admin,admin_id),(student,student_id),(outsider,stranger_id)):
                with client.session_transaction() as sess:
                    sess["user_id"] = uid
                    sess["user_name"] = "Test User"
                    if uid == admin_id:
                        sess["admin_csrf_token"]="test-preview-token"
            assert outsider.get(f"/subjects/{sid}").status_code == 302

            # Admin creates the official source, student does not receive edit rights.
            shared = project.ensure_shared_subject_for_user_subject(conn,student_id,sid)
            shared_id = shared["id"]
            conn.execute("UPDATE shared_subjects SET admin_managed=1 WHERE id=?",(shared_id,))
            conn.commit()
            uploaded = admin.post("/admin/documents",data={
                "csrf_token":"test-preview-token",
                "shared_subject_id":str(shared_id), "is_syllabus":"1",
                "file":(io.BytesIO(src.read_bytes()),"de_cuong.docx")
            },content_type="multipart/form-data")
            assert uploaded.status_code in (302,303),uploaded.status_code
            row=conn.execute("""
                SELECT id,extracted_text,stored_name FROM shared_documents
                WHERE shared_subject_id=? AND is_syllabus=1
            """,(shared_id,)).fetchone()
            assert row and "Chương 1" in row["extracted_text"]
            doc_id=row["id"]
            file_path=f"/shared-documents/{doc_id}/file"
            for client in (student,admin):
                response=client.get(file_path+"?preview=1")
                assert response.status_code==200
                assert "text/html" in response.content_type
                assert "frame-ancestors 'self'" in response.headers["Content-Security-Policy"]
                assert "<table>" in response.get_data(as_text=True)
                assert "<script>" not in response.get_data(as_text=True)
                download=client.get(file_path+"?download=1")
                assert download.status_code==200
                assert "attachment" in download.headers["Content-Disposition"]
            assert outsider.get(file_path+"?preview=1").status_code==403
            assert outsider.get(file_path+"?download=1").status_code==403

            from chapter_material_routes import chapter_key
            from syllabus_routes import ensure_syllabus_schema
            ensure_syllabus_schema(conn)
            ensure_chapter_material_schema(conn)
            units=[
                {"title":"Chương 1. Sự ra đời của Đảng",
                 "evidence":"Chương 1. Sự ra đời của Đảng"},
                {"title":"Chương 2. Lãnh đạo kháng chiến",
                 "evidence":"Chương 2. Lãnh đạo kháng chiến"},
            ]
            import json
            fingerprint=syllabus_hash(row["extracted_text"])
            conn.execute("""
                INSERT INTO syllabus_reviews(user_id,subject_id,source_hash,source_kind,
                                            units_json,status,updated_at)
                VALUES(?,?,?,?,?,'confirmed',?)
            """,(student_id,sid,fingerprint,"shared",json.dumps(units),stamp))
            slide=conn.execute("""
                INSERT INTO shared_documents(shared_subject_id,original_name,
                        stored_name,extracted_text,is_syllabus,created_by_user_id,created_at)
                VALUES(?,'chuong1.pdf','chuong1-test.pdf','Lịch sử chương 1',0,?,?)
            """,(shared_id,admin_id,stamp)).lastrowid
            writer=PdfWriter()
            writer.add_blank_page(width=595,height=842)
            with open(project.UPLOAD_DIR/"chuong1-test.pdf","wb") as fh:
                writer.write(fh)
            conn.execute("""
                INSERT INTO shared_document_chapters(document_id,shared_subject_id,
                    chapter_key,syllabus_hash,confirmed_by,confirmed_at)
                VALUES(?,?,?,?,?,?)
            """,(slide,shared_id,"1",fingerprint,admin_id,stamp))
            conn.commit()

            page=student.get(f"/subjects/{sid}")
            assert page.status_code==200, page.get_data(as_text=True)[:650]
            html=page.get_data(as_text=True)
            assert "course-material-split" in html
            assert "course_overview_refresh.css" in html
            assert "Xem file chi tiết" in html
            assert "Tài liệu theo chương" in html
            assert "1 tài liệu được duyệt" in html
            assert "Chưa có tài liệu chung được duyệt" in html
            assert "Cài đặt học tập" in html
            assert 'id="courseSettings"' in html
            assert 'class="shared-syllabus-readonly"' not in html
            assert "Đề cương → Phân tích chương" not in html
            assert "Hỏi Study Agent" not in html
            assert "Chưa có mô tả cho môn" not in html
            assert html.count('class="course-chapter-item"') == 2
            assert 'href="/subjects/' + str(sid) + '/chapters/1"' in html
            assert student.get(f"/subjects/{sid}/chapters/1").status_code==200
            assert student.get(f"/subjects/{sid}/chapters/2").status_code==200

            # Confirm preference edits remain persistent inside collapsed form.
            profile=conn.execute("SELECT notes FROM learning_profiles WHERE user_id=?",
                                 (student_id,)).fetchone()
            assert profile
            saved=student.post(f"/subjects/{sid}",data={
                "action_type":"profile", "learning_style":"visual",
                "preferred_study_time":"20:30", "daily_minutes":"60",
                "goals":"Hoc chuong 1","notes":"Ghi chu cua toi"
            })
            assert saved.status_code==302
            prefs=conn.execute("SELECT * FROM learning_profiles WHERE user_id=?",
                               (student_id,)).fetchone()
            assert prefs["notes"]=="Ghi chu cua toi"
            assert prefs["daily_minutes"]==60
            assert "20:30" in student.get(f"/subjects/{sid}").get_data(as_text=True)

            # A PDF syllabus has a native inline viewer; uploaded Word uses local HTML.
            pdf_doc=conn.execute("""
                INSERT INTO shared_documents(shared_subject_id,original_name,
                        stored_name,extracted_text,is_syllabus,created_by_user_id,created_at)
                VALUES(?,'pdf-outline.pdf','pdf-outline.pdf','Chương 1. PDF',1,?,?)
            """,(shared_id,admin_id,stamp)).lastrowid
            conn.execute("UPDATE shared_documents SET is_syllabus=0 WHERE id=?",(doc_id,))
            conn.execute("UPDATE shared_subjects SET syllabus_text=? WHERE id=?",
                         ("Chương 1. PDF",shared_id))
            (project.UPLOAD_DIR/"pdf-outline.pdf").write_bytes(
                (project.UPLOAD_DIR/"chuong1-test.pdf").read_bytes())
            conn.commit()
            pdf_preview=student.get(f"/shared-documents/{pdf_doc}/file?preview=1")
            assert pdf_preview.status_code==200
            assert "application/pdf" in pdf_preview.content_type
            pdf_page=student.get(f"/subjects/{sid}").get_data(as_text=True)
            assert "pdf-outline.pdf" in pdf_page and "?preview=1" in pdf_page
            conn.close()
        finally:
            project.DB_PATH,project.UPLOAD_DIR=old_db,old_upload
    print("Course refresh QA PASS: DOCX/PDF preview, secure ownership, confirmed chapters, 50/50 layout")


if __name__=="__main__":
    run()
